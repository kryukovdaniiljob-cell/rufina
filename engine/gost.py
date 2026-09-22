# -*- coding: utf-8 -*-
"""Вёрстка документов по ГОСТ 7.32-2017.

Один слой на все текстовые отчёты: параметры набора, заголовки, таблицы,
перечисления и нумерация живут здесь, а не расползаются по модулям.
Правила взяты из стандарта: А4, поля 30/15/20/20 мм, Times New Roman
14 пт, полуторный интервал, абзацный отступ 1,25 см, выравнивание по
ширине; полужирный только в заголовках; таблицы нумеруются сквозной
нумерацией, название ставится слева над таблицей.
"""
from __future__ import annotations
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as ВЫР
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ОСН = "Times New Roman"
КЕГЛЬ = Pt(14)
КЕГЛЬ_ТАБ = Pt(12)          # в таблице допускается шрифт меньшего размера
ОТСТУП = Cm(1.25)


def _шрифт(элемент, имя=ОСН):
    """Кириллице нужен явный шрифт во всех слотах, иначе Word подставит свой."""
    rpr = элемент.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts")
        rpr.append(f)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        f.set(qn(a), имя)


def _уровень(параграф, лвл):
    """Уровень структуры: без него заголовок не попадёт в оглавление."""
    ppr = параграф._p.get_or_add_pPr()
    o = ppr.find(qn("w:outlineLvl"))
    if o is None:
        o = OxmlElement("w:outlineLvl")
        ppr.append(o)
    o.set(qn("w:val"), str(лвл))


class Документ:
    """Документ по ГОСТ. Нумерация разделов и таблиц ведётся сама."""

    def __init__(self):
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin, s.right_margin = Mm(30), Mm(15)
        s.top_margin, s.bottom_margin = Mm(20), Mm(20)
        s.different_first_page_header_footer = True     # титул без номера

        n = self.doc.styles["Normal"]
        n.font.name = ОСН
        n.font.size = КЕГЛЬ
        n.font.color.rgb = RGBColor.from_string("000000")
        _шрифт(n.element)
        p = n.paragraph_format
        p.line_spacing = 1.5
        p.first_line_indent = ОТСТУП
        p.alignment = ВЫР.JUSTIFY
        p.space_before = Pt(0)
        p.space_after = Pt(0)

        self._раздел = 0
        self._подраздел = 0
        self._таблица = 0
        self._номер_страницы(s)

    # ------------------------------------------------------ служебное
    def _номер_страницы(self, section):
        p = section.footer.paragraphs[0]
        p.alignment = ВЫР.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run()
        r.font.name, r.font.size = ОСН, КЕГЛЬ
        _шрифт(r._element)
        for тип, текст in (("begin", None), (None, "PAGE"), ("end", None)):
            if тип:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), тип)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = текст
            r._element.append(e)

    def _прогон(self, параграф, текст, жирный=False, кегль=None):
        r = параграф.add_run(str(текст))
        r.bold = жирный
        r.font.size = кегль or КЕГЛЬ
        r.font.name = ОСН
        r.font.color.rgb = RGBColor.from_string("000000")
        _шрифт(r._element)
        return r

    # ------------------------------------------------------ текст
    def абзац(self, текст="", жирный=False, выравн=None, отступ=True,
              перед=0, после=0, разрыв=False, кегль=None):
        p = self.doc.add_paragraph()
        f = p.paragraph_format
        f.page_break_before = разрыв
        f.first_line_indent = ОТСТУП if отступ else Pt(0)
        f.alignment = выравн if выравн is not None else ВЫР.JUSTIFY
        f.space_before, f.space_after = Pt(перед), Pt(после)
        f.line_spacing = 1.5
        if текст:
            self._прогон(p, текст, жирный, кегль)
        return p

    def пусто(self, сколько=1):
        for _ in range(сколько):
            self.абзац("")

    def структурный(self, текст):
        """РЕФЕРАТ, СОДЕРЖАНИЕ, ВВЕДЕНИЕ, ЗАКЛЮЧЕНИЕ — прописными по центру."""
        p = self.абзац(текст.upper(), жирный=True, выравн=ВЫР.CENTER,
                       отступ=False, после=18, разрыв=True)
        _уровень(p, 0)
        self._раздел, self._подраздел = self._раздел, 0
        return p

    def раздел(self, текст):
        """Нумерованный раздел с абзацного отступа, полужирным, с новой страницы."""
        self._раздел += 1
        self._подраздел = 0
        p = self.абзац("%d %s" % (self._раздел, текст), жирный=True,
                       выравн=ВЫР.LEFT, после=12, разрыв=True)
        _уровень(p, 0)
        return p

    def подраздел(self, текст):
        self._подраздел += 1
        p = self.абзац("%d.%d %s" % (self._раздел, self._подраздел, текст),
                       жирный=True, выравн=ВЫР.LEFT, перед=12, после=6)
        _уровень(p, 1)
        return p

    def список(self, пункты, маркер="- "):
        """Перечисление через тире с абзацного отступа [1, п. 6.4.6]."""
        for п in пункты:
            self.абзац(маркер + str(п))

    # ------------------------------------------------------ таблицы
    def таблица(self, название, шапка, строки, ширины=None, право=(),
                ссылка=True):
        """«Таблица N — Название» слева над таблицей, ссылка в тексте."""
        self._таблица += 1
        н = self._таблица
        if ссылка:
            фраза = ("Данные приведены в таблице %d." % н if н % 2
                     else "Показатели сведены в таблице %d." % н)
            self.абзац(фраза)
        self.абзац("Таблица %d — %s" % (н, название), отступ=False,
                   выравн=ВЫР.LEFT, перед=6, после=2)
        t = self.doc.add_table(rows=1, cols=len(шапка))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        def яч(c, v, шап, вправо):
            c.text = ""
            p = c.paragraphs[0]
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.line_spacing = 1.0
            pf.space_before, pf.space_after = Pt(2), Pt(2)
            p.alignment = (ВЫР.CENTER if шап else
                           ВЫР.RIGHT if вправо else ВЫР.LEFT)
            self._прогон(p, v, шап, КЕГЛЬ_ТАБ)

        for j, h in enumerate(шапка):
            яч(t.rows[0].cells[j], h, True, False)
        for стр in строки:
            cells = t.add_row().cells
            for j, v in enumerate(стр):
                яч(cells[j], "" if v is None else v, False, j in право)
        if ширины:
            for j, w in enumerate(ширины):
                for rw in t.rows:
                    rw.cells[j].width = Mm(w)
        хв = self.абзац("", после=6)
        хв.paragraph_format.line_spacing = 1.0
        return t

    # ------------------------------------------------------ поля Word
    def оглавление(self):
        p = self.абзац("", отступ=False, после=0)
        r = p.add_run()
        r.font.name, r.font.size = ОСН, КЕГЛЬ
        _шрифт(r._element)
        части = [("begin", None), (None, r'TOC \o "1-2" \h \z \u'),
                 ("separate", None), ("текст", "Оглавление собирается при "
                                                "открытии файла (F9)."),
                 ("end", None)]
        for тип, знач in части:
            if тип == "текст":
                e = OxmlElement("w:t")
                e.text = знач
            elif тип is None:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = знач
            else:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), тип)
            r._element.append(e)
        return p

    def сохранить(self, путь):
        """Поля пересчитываются при открытии: оглавление и номера страниц."""
        s = self.doc.settings.element
        u = s.find(qn("w:updateFields"))
        if u is None:
            u = OxmlElement("w:updateFields")
            s.append(u)
        u.set(qn("w:val"), "true")
        self.doc.save(путь)
        return путь
